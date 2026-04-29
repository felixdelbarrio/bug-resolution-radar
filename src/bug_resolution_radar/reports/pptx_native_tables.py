"""Native PowerPoint table helpers for executive reports."""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Any, Mapping, Sequence

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt


@dataclass(frozen=True)
class NativeTableColumnSpec:
    header: str
    width: int
    align: str = "left"
    font_size_pt: float = 9.5
    header_font_size_pt: float = 8.5
    bold: bool = False
    max_chars: int | None = None


@dataclass(frozen=True)
class NativeTableStyle:
    header_fill: RGBColor
    header_font: RGBColor
    body_font: RGBColor
    border: RGBColor
    zebra_fill: RGBColor
    body_fill: RGBColor
    font_name: str = "Arial"


DEFAULT_NATIVE_TABLE_STYLE = NativeTableStyle(
    header_fill=RGBColor(4, 19, 139),
    header_font=RGBColor(255, 255, 255),
    body_font=RGBColor(4, 19, 139),
    border=RGBColor(205, 214, 232),
    zebra_fill=RGBColor(248, 250, 255),
    body_fill=RGBColor(255, 255, 255),
    font_name="Arial",
)

_TABLE_BORDER_WIDTH_EMU = 10_160
_CELL_MARGIN_LEFT = int(Pt(3.8))
_CELL_MARGIN_RIGHT = int(Pt(3.2))
_CELL_MARGIN_TOP = int(Pt(1.4))
_CELL_MARGIN_BOTTOM = int(Pt(1.4))


def ellipsize_text(value: object, *, max_chars: int) -> str:
    """Trim text without splitting words when there is a sensible break point."""
    text = re.sub(r"\s+", " ", str(value or "").strip())
    if max_chars <= 0 or len(text) <= max_chars:
        return text
    if max_chars == 1:
        return "…"

    limit = max(int(max_chars) - 1, 1)
    candidate = text[:limit].rstrip()
    min_break = max(int(limit * 0.62), 1)
    break_positions = [
        candidate.rfind(separator) for separator in (" ", "/", "-", ":", ";", ",", ".")
    ]
    best_break = max(break_positions or [-1])
    if best_break >= min_break:
        candidate = candidate[:best_break].rstrip(" /-:;,.")
    if not candidate:
        candidate = text[:limit].rstrip()
    return f"{candidate}…"


def native_column_widths(total_width: int, weights: Sequence[float | int]) -> tuple[int, ...]:
    """Convert relative weights into deterministic EMU column widths."""
    safe_total = max(int(total_width or 0), 1)
    raw = [max(float(value or 0), 0.0) for value in list(weights or [])]
    if not raw:
        return ()
    if sum(raw) <= 0.0:
        raw = [1.0 for _ in raw]
    weight_total = sum(raw)
    widths = [max(int(round((value / weight_total) * float(safe_total))), 1) for value in raw]
    widths[-1] += safe_total - sum(widths)
    return tuple(widths)


def rebuild_native_table_shape(
    slide: Any,
    table_shape: Any | None,
    *,
    rows: int,
    cols: int,
    geometry: tuple[int, int, int, int] | None = None,
) -> Any:
    """Replace an existing table placeholder with a native table of fixed size."""
    left, top, width, height = 0, 0, 1, 1
    if table_shape is not None:
        left = int(getattr(table_shape, "left", 0) or 0)
        top = int(getattr(table_shape, "top", 0) or 0)
        width = max(int(getattr(table_shape, "width", 0) or 0), 1)
        height = max(int(getattr(table_shape, "height", 0) or 0), 1)
    if geometry is not None:
        g_left, g_top, g_width, g_height = geometry
        left = int(g_left)
        top = int(g_top)
        width = max(int(g_width), 1)
        height = max(int(g_height), 1)

    if table_shape is not None:
        try:
            node = table_shape.element
            node.getparent().remove(node)
        except Exception:
            pass
    return slide.shapes.add_table(
        max(int(rows or 0), 1),
        max(int(cols or 0), 1),
        left,
        top,
        width,
        height,
    )


def populate_native_table(
    table_shape: Any,
    *,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    column_widths: Sequence[int],
    row_height: int,
    header_height: int,
    font_name: str,
    body_font_size_pt: float,
    header_font_size_pt: float,
    left_align_cols: Sequence[int] = (),
    center_align_cols: Sequence[int] = (),
    hyperlink_by_row: Mapping[int, str] | None = None,
    zebra: bool = True,
) -> None:
    """Populate and style an existing native PowerPoint table."""
    if table_shape is None or not getattr(table_shape, "has_table", False):
        return

    table = table_shape.table
    col_count = len(table.columns)
    if col_count <= 0:
        return
    row_count = len(table.rows)
    if row_count <= 0:
        return

    normalized_headers = [str(value or "") for value in list(headers or [])[:col_count]]
    normalized_headers.extend([""] * max(col_count - len(normalized_headers), 0))
    normalized_rows = [list(row)[:col_count] for row in list(rows or [])]
    if not normalized_rows and row_count > 1:
        normalized_rows = [[""] * col_count]
    normalized_rows = [
        [str(value or "") for value in row] + ([""] * max(col_count - len(row), 0))
        for row in normalized_rows[: max(row_count - 1, 0)]
    ]

    widths = [int(width or 0) for width in list(column_widths or [])[:col_count]]
    if len(widths) < col_count:
        remaining = max(int(getattr(table_shape, "width", 0) or 0) - sum(widths), 1)
        fallback = max(int(round(remaining / max(col_count - len(widths), 1))), 1)
        widths.extend([fallback] * (col_count - len(widths)))
    widths[-1] += max(int(getattr(table_shape, "width", 0) or 0) - sum(widths), 0)
    for col, width in zip(table.columns, widths):
        try:
            col.width = max(int(width), 1)
        except Exception:
            pass

    try:
        table.rows[0].height = max(int(header_height or 0), 1)
    except Exception:
        pass
    for ridx in range(1, row_count):
        try:
            table.rows[ridx].height = max(int(row_height or 0), 1)
        except Exception:
            pass

    left_cols = {int(idx) for idx in left_align_cols}
    center_cols = {int(idx) for idx in center_align_cols}
    links = {
        int(key): str(value).strip()
        for key, value in dict(hyperlink_by_row or {}).items()
        if str(value or "").strip()
    }
    style = DEFAULT_NATIVE_TABLE_STYLE
    safe_font_name = str(font_name or style.font_name).strip() or style.font_name

    for cidx in range(col_count):
        cell = table.cell(0, cidx)
        _apply_fill(cell, style.header_fill)
        _set_cell_text(
            cell,
            normalized_headers[cidx],
            align=_column_alignment(cidx, left_cols=left_cols, center_cols=center_cols),
            font_name=safe_font_name,
            font_size_pt=float(header_font_size_pt),
            color_rgb=style.header_font,
            bold=True,
            hyperlink="",
            word_wrap=True,
        )
        _set_cell_border(cell, color_rgb=style.border, width_emu=_TABLE_BORDER_WIDTH_EMU)

    for ridx in range(1, row_count):
        row_values = (
            normalized_rows[ridx - 1] if ridx - 1 < len(normalized_rows) else [""] * col_count
        )
        fill_rgb = style.zebra_fill if zebra and ridx % 2 == 0 else style.body_fill
        for cidx in range(col_count):
            cell = table.cell(ridx, cidx)
            _apply_fill(cell, fill_rgb)
            _set_cell_text(
                cell,
                row_values[cidx] if cidx < len(row_values) else "",
                align=_column_alignment(cidx, left_cols=left_cols, center_cols=center_cols),
                font_name=safe_font_name,
                font_size_pt=float(body_font_size_pt),
                color_rgb=style.body_font,
                bold=False,
                hyperlink=links.get(ridx - 1, "") if cidx == 0 else "",
                word_wrap=True,
            )
            _set_cell_border(cell, color_rgb=style.border, width_emu=_TABLE_BORDER_WIDTH_EMU)


def _column_alignment(
    cidx: int,
    *,
    left_cols: set[int],
    center_cols: set[int],
) -> PP_ALIGN:
    if cidx in center_cols:
        return PP_ALIGN.CENTER
    if cidx in left_cols:
        return PP_ALIGN.LEFT
    return PP_ALIGN.LEFT


def _apply_fill(cell: Any, color_rgb: RGBColor) -> None:
    try:
        cell.fill.solid()
        cell.fill.fore_color.rgb = color_rgb
    except Exception:
        pass


def _set_cell_text(
    cell: Any,
    text: str,
    *,
    align: PP_ALIGN,
    font_name: str,
    font_size_pt: float,
    color_rgb: RGBColor,
    bold: bool,
    hyperlink: str,
    word_wrap: bool,
) -> None:
    tf = getattr(cell, "text_frame", None)
    if tf is None:
        return
    try:
        tf.clear()
    except Exception:
        pass
    try:
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.word_wrap = bool(word_wrap)
        tf.margin_left = _CELL_MARGIN_LEFT
        tf.margin_right = _CELL_MARGIN_RIGHT
        tf.margin_top = _CELL_MARGIN_TOP
        tf.margin_bottom = _CELL_MARGIN_BOTTOM
    except Exception:
        pass
    try:
        cell.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    except Exception:
        pass

    paragraphs = list(tf.paragraphs)
    if not paragraphs:
        paragraph = tf.add_paragraph()
    else:
        paragraph = paragraphs[0]
    try:
        paragraph.clear()
    except Exception:
        pass
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)

    run = paragraph.add_run()
    run.text = str(text or "")
    run.font.size = Pt(max(float(font_size_pt), 1.0))
    try:
        run.font.name = font_name
    except Exception:
        pass
    run.font.bold = bool(bold)
    try:
        run.font.color.rgb = color_rgb
    except Exception:
        pass
    if hyperlink:
        try:
            run.hyperlink.address = str(hyperlink).strip()
            run.font.underline = True
        except Exception:
            pass


def _set_cell_border(cell: Any, *, color_rgb: RGBColor, width_emu: int) -> None:
    tc = getattr(cell, "_tc", None)
    if tc is None:
        return
    tc_pr = tc.get_or_add_tcPr()
    color_hex = f"{int(color_rgb[0]):02X}{int(color_rgb[1]):02X}{int(color_rgb[2]):02X}"
    for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        existing = tc_pr.find(qn(side))
        if existing is not None:
            tc_pr.remove(existing)
        ln = OxmlElement(side)
        ln.set("w", str(max(int(width_emu or 0), 1)))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")

        solid = OxmlElement("a:solidFill")
        srgb = OxmlElement("a:srgbClr")
        srgb.set("val", color_hex)
        solid.append(srgb)
        ln.append(solid)

        dash = OxmlElement("a:prstDash")
        dash.set("val", "solid")
        ln.append(dash)

        head = OxmlElement("a:headEnd")
        head.set("type", "none")
        head.set("w", "med")
        head.set("len", "med")
        ln.append(head)

        tail = OxmlElement("a:tailEnd")
        tail.set("type", "none")
        tail.set("w", "med")
        tail.set("len", "med")
        ln.append(tail)
        tc_pr.append(ln)
